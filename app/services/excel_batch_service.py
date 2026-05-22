import os
import time
from typing import List

from flask import current_app
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from app import db
from app.services.arquivo_service import ArquivoService


class ExcelBatchService:
    ALLOWED_EXTENSIONS = {"xlsx"}
    BATCH_CATEGORY = "DOCUMENTO"
    BATCH_SUBFOLDER = "batch"

    @classmethod
    def is_allowed_file(cls, filename: str) -> bool:
        if not filename:
            return False
        filename = os.path.basename(filename)
        if filename.startswith("~$"):
            return False
        extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        return extension in cls.ALLOWED_EXTENSIONS

    @classmethod
    def normalize_filename(cls, filename: str) -> str:
        filename = os.path.basename(filename or "")
        return secure_filename(filename)

    @classmethod
    def process_batch(cls, files: List, empresa_id: int, usuario_id: int) -> dict:
        start_at = time.time()
        current_app.logger.info(
            "Batch upload XLSX iniciado",
            extra={"empresa_id": empresa_id, "usuario_id": usuario_id, "total_arquivos": len(files)},
        )

        resultados = []
        sucesso = 0
        falhas = 0

        for arquivo in files:
            nome_arquivo = getattr(arquivo, "filename", "<desconhecido>")
            try:
                resultado = cls.process_single(arquivo, empresa_id, usuario_id)
                resultados.append({"arquivo": nome_arquivo, "sucesso": True, "detalhes": resultado})
                sucesso += 1
            except Exception as exc:
                current_app.logger.warning(
                    "Falha no processamento de arquivo batch",
                    extra={"arquivo": nome_arquivo, "erro": str(exc)},
                )
                resultados.append({"arquivo": nome_arquivo, "sucesso": False, "erro": str(exc)})
                falhas += 1

        tempo_total = round(time.time() - start_at, 3)
        current_app.logger.info(
            "Batch upload XLSX finalizado",
            extra={
                "empresa_id": empresa_id,
                "usuario_id": usuario_id,
                "sucesso": sucesso,
                "falhas": falhas,
                "tempo_total_segundos": tempo_total,
            },
        )

        return {
            "arquivos_processados": len(files),
            "sucesso": sucesso,
            "falhas": falhas,
            "tempo_total_segundos": tempo_total,
            "detalhes": resultados,
        }

    @classmethod
    def process_single(cls, file_storage, empresa_id: int, usuario_id: int) -> dict:
        filename = cls.normalize_filename(getattr(file_storage, "filename", ""))
        if not cls.is_allowed_file(filename):
            raise ValueError("Apenas arquivos .xlsx válidos podem ser processados.")

        file_storage.stream.seek(0)

        try:
            workbook = load_workbook(file_storage.stream, data_only=True)
        except Exception as exc:
            raise ValueError(f"Erro ao ler arquivo Excel: {exc}")

        if not workbook.sheetnames:
            raise ValueError("O arquivo Excel não contém nenhuma planilha.")

        worksheet = workbook.active
        if worksheet.max_row == 0 and worksheet.max_column == 0:
            raise ValueError("A planilha selecionada está vazia.")

        primeira_linha = [cell for cell in next(worksheet.iter_rows(max_row=1, values_only=True))] if worksheet.max_row else []
        summary = {
            "planilha": worksheet.title,
            "linhas": worksheet.max_row,
            "colunas": worksheet.max_column,
            "cabecalho": [str(value) if value is not None else "" for value in primeira_linha],
        }

        file_storage.stream.seek(0)
        arquivo_registrado = ArquivoService.save_local_file(
            file_storage=file_storage,
            empresa_id=empresa_id,
            usuario_id=usuario_id,
            categoria=cls.BATCH_CATEGORY,
            entidade="BATCH_UPLOAD",
            entidade_id=None,
            obra_id=None,
            rdo_id=None,
            publico=False,
            subfolder=cls.BATCH_SUBFOLDER,
        )

        db.session.commit()

        summary["arquivo_id"] = arquivo_registrado.id
        summary["nome_original"] = arquivo_registrado.nome_original
        summary["storage_path"] = arquivo_registrado.storage_path

        return summary
